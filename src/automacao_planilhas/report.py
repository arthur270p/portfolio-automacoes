"""Resumo da execução e geração do relatório Excel.

O relatório é o produto entregue: quem abre não vê o código, vê estas três
abas. Por isso data e decimal saem como valor do Excel, e não como texto —
texto não soma nem ordena, e uma planilha que não soma não resolve o problema
de ninguém.
"""

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .models import ReadBatch, ValidationResult

SHEET_CLEAN = "dados_limpos"
SHEET_ERRORS = "erros"
SHEET_SUMMARY = "resumo"

SUMMARY_COLUMNS = ("categoria", "item", "valor", "detalhe")

_DATE_FORMAT = "DD/MM/YYYY"
_DECIMAL_FORMAT = "#,##0.00"
_MAX_COLUMN_WIDTH = 60
_MIN_COLUMN_WIDTH = 10


def build_summary(batch: ReadBatch, validation: ValidationResult) -> pd.DataFrame:
    """Monta o resumo: sete métricas e uma linha por origem com ocorrência.

    As origens entram nominalmente, e não como contagem, porque consolidar
    doze arquivos e receber nove sem perceber é o pesadelo desse tipo de
    ferramenta. Quem leu o resumo precisa conseguir dizer qual arquivo ficou
    de fora e por quê.
    """
    rows: list[dict[str, object]] = [
        {"categoria": "metrica", "item": item, "valor": valor, "detalhe": detalhe}
        for item, valor, detalhe in (
            ("arquivos_encontrados", batch.files_found, "Arquivos CSV ou XLSX na pasta de entrada."),
            ("arquivos_processados", batch.files_processed, "Arquivos aceitos e consolidados."),
            ("registros_lidos", int(len(batch.data)), "Linhas lidas das origens aceitas."),
            ("registros_validos", int(len(validation.clean_data)), "Linhas sem nenhum problema."),
            ("registros_invalidos", int(len(validation.error_data)), "Linhas com ao menos um problema."),
            ("duplicidades", validation.duplicate_count, "Repetições da chave; a primeira foi mantida."),
            ("problemas_encontrados", validation.problem_count, "Total de problemas, somando os de uma mesma linha."),
        )
    ]

    for issue in batch.source_issues:
        rows.append(
            {
                "categoria": "origem",
                "item": f"{issue.file_name} ({issue.sheet_name})" if issue.sheet_name else issue.file_name,
                "valor": issue.level,
                "detalhe": f"{issue.code}: {issue.detail}",
            }
        )

    return pd.DataFrame(rows, columns=list(SUMMARY_COLUMNS))


def _neutralize_and_format(sheet) -> None:
    """Corrige tipo e formato célula a célula, depois da escrita do pandas.

    Três ajustes, cada um medido contra o comportamento real do openpyxl:

    1. Texto iniciado por `=` é gravado como **fórmula**, o que faria o Excel
       executar o conteúdo da planilha de origem ao abrir o relatório. Medido:
       no XLSX apenas o `=` dispara isso — `+`, `-` e `@` ficam como texto.
       Tratar os quatro corromperia valores legítimos como `-5`.
    2. Data sai no formato ISO do openpyxl; quem vai ler espera dd/mm/aaaa.
    3. Decimal sai sem formatação. Inteiro fica como está: aplicar duas casas
       a uma contagem transformaria `2` em `2,00`.
    """
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            value = cell.value
            if cell.data_type == "f" and isinstance(value, str):
                cell.data_type = "s"
            elif isinstance(value, (datetime, date)):
                cell.number_format = _DATE_FORMAT
            elif isinstance(value, (float, Decimal)):
                # `Decimal` ainda nao virou `float` neste ponto: a conversao
                # acontece ao salvar. Testar so por `float` perderia
                # exatamente os valores monetarios que o conversor produz.
                cell.number_format = _DECIMAL_FORMAT


def _finish_sheet(sheet) -> None:
    if sheet.max_row == 0 or sheet.max_column == 0:
        return

    for cell in sheet[1]:
        cell.font = Font(bold=True)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(sheet.max_column)}{max(sheet.max_row, 1)}"
    )

    for index in range(1, sheet.max_column + 1):
        longest = max(
            (len(str(cell.value)) for cell in sheet[get_column_letter(index)] if cell.value is not None),
            default=0,
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(longest + 2, _MIN_COLUMN_WIDTH), _MAX_COLUMN_WIDTH
        )

    _neutralize_and_format(sheet)


def write_report(
    path: Path,
    clean_data: pd.DataFrame,
    error_data: pd.DataFrame,
    summary: pd.DataFrame,
) -> None:
    """Grava as três abas no caminho informado.

    A gravação atômica — arquivo temporário e troca no fim — é da CLI, que é
    quem conhece o destino final e pode limpar o rastro em caso de falha.
    Aqui o contrato é simples: ou o arquivo sai completo, ou a exceção sobe.
    """
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        clean_data.to_excel(writer, sheet_name=SHEET_CLEAN, index=False)
        error_data.to_excel(writer, sheet_name=SHEET_ERRORS, index=False)
        summary.to_excel(writer, sheet_name=SHEET_SUMMARY, index=False)

        for name in (SHEET_CLEAN, SHEET_ERRORS, SHEET_SUMMARY):
            _finish_sheet(writer.book[name])
