import io
import json
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import mock

import openpyxl

from automacao_planilhas import __version__
from automacao_planilhas.cli import main

VALID_CSV = "nome;email;data;valor\nAna;ana@example.com;10/09/2026;12,50\n"

CONFIG = {
    "colunas_obrigatorias": ["nome", "email", "data", "valor"],
    "chaves_duplicidade": ["email", "data"],
    "tipos": {"data": "data_br", "valor": "decimal_br"},
}


class Workspace:
    """Pasta de trabalho pronta: entrada, configuração e destino."""

    def __init__(self, root: Path, csv_text: str = VALID_CSV, config: dict | None = None):
        self.root = root
        self.input_dir = root / "entrada"
        self.input_dir.mkdir()
        if csv_text:
            (self.input_dir / "dados.csv").write_text(csv_text, encoding="utf-8-sig")
        self.config = root / "config.json"
        self.config.write_text(
            json.dumps(config if config is not None else CONFIG), encoding="utf-8"
        )
        self.output = root / "saida" / "relatorio.xlsx"

    def argv(self, *extra: str) -> list[str]:
        return [
            "processar",
            "--entrada",
            str(self.input_dir),
            "--config",
            str(self.config),
            "--saida",
            str(self.output),
            *extra,
        ]


def write_garbage_then_fail(path, *_args, **_kwargs):
    """Escreve no caminho recebido e só então falha.

    Falhar antes de escrever nao prova nada: o destino ficaria intacto com ou
    sem arquivo temporario. O risco real e a escrita que comeca, suja o
    caminho e morre no meio.
    """
    Path(path).write_bytes(b"conteudo pela metade")
    raise OSError("disco cheio")


class CliTests(unittest.TestCase):
    def test_processes_example_and_protects_existing_output(self):
        with TemporaryDirectory() as temporary:
            workspace = Workspace(Path(temporary))

            first = main(workspace.argv())
            second = main(workspace.argv())

            self.assertEqual(first, 0)
            # Sobrescrever um relatorio sem pedir apaga trabalho de alguem.
            self.assertEqual(second, 4)
            self.assertTrue(workspace.output.exists())

    def test_overwrite_flag_replaces_the_existing_report(self):
        with TemporaryDirectory() as temporary:
            workspace = Workspace(Path(temporary))

            self.assertEqual(main(workspace.argv()), 0)
            self.assertEqual(main(workspace.argv("--sobrescrever")), 0)

            workbook = openpyxl.load_workbook(workspace.output)
            try:
                self.assertEqual(
                    workbook.sheetnames, ["dados_limpos", "erros", "resumo"]
                )
            finally:
                workbook.close()

    def test_creates_the_output_directory(self):
        with TemporaryDirectory() as temporary:
            workspace = Workspace(Path(temporary))
            self.assertFalse(workspace.output.parent.exists())

            self.assertEqual(main(workspace.argv()), 0)
            self.assertTrue(workspace.output.exists())

    def test_invalid_configuration_stops_before_reading_data(self):
        with TemporaryDirectory() as temporary:
            workspace = Workspace(
                Path(temporary), config={"colunas_obrigatorias": ["Nome Errado"]}
            )

            self.assertEqual(main(workspace.argv()), 2)
            self.assertFalse(workspace.output.exists())

    def test_missing_input_directory_is_an_argument_error(self):
        with TemporaryDirectory() as temporary:
            workspace = Workspace(Path(temporary))
            argv = workspace.argv()
            argv[argv.index("--entrada") + 1] = str(Path(temporary) / "ausente")

            self.assertEqual(main(argv), 2)

    def test_empty_input_reports_no_valid_source(self):
        with TemporaryDirectory() as temporary:
            workspace = Workspace(Path(temporary), csv_text="")

            self.assertEqual(main(workspace.argv()), 3)
            self.assertFalse(workspace.output.exists())

    def test_all_sources_rejected_reports_no_valid_source(self):
        with TemporaryDirectory() as temporary:
            # Coluna obrigatoria ausente: a origem inteira e recusada.
            workspace = Workspace(Path(temporary), csv_text="nome\nAna\n")

            self.assertEqual(main(workspace.argv()), 3)
            self.assertFalse(workspace.output.exists())

    def test_failure_while_writing_leaves_no_temporary_file_behind(self):
        with TemporaryDirectory() as temporary:
            workspace = Workspace(Path(temporary))

            with mock.patch(
                "automacao_planilhas.cli.write_report",
                side_effect=write_garbage_then_fail,
            ):
                self.assertEqual(main(workspace.argv()), 1)

            self.assertFalse(workspace.output.exists())
            leftovers = list(workspace.output.parent.glob("*.tmp.xlsx"))
            self.assertEqual(leftovers, [])

    def test_failure_while_writing_preserves_the_previous_report(self):
        # O relatorio anterior e a unica copia boa que a pessoa tem; uma falha
        # no meio da escrita nao pode destrui-la.
        with TemporaryDirectory() as temporary:
            workspace = Workspace(Path(temporary))
            self.assertEqual(main(workspace.argv()), 0)
            original = workspace.output.read_bytes()

            with mock.patch(
                "automacao_planilhas.cli.write_report",
                side_effect=write_garbage_then_fail,
            ):
                self.assertEqual(main(workspace.argv("--sobrescrever")), 1)

            self.assertEqual(workspace.output.read_bytes(), original)

    def test_reports_invalid_rows_without_failing_the_run(self):
        with TemporaryDirectory() as temporary:
            workspace = Workspace(
                Path(temporary),
                csv_text=(
                    "nome;email;data;valor\n"
                    "Ana;ana@example.com;10/09/2026;12,50\n"
                    ";sem-nome@example.com;11/09/2026;9,90\n"
                ),
            )

            self.assertEqual(main(workspace.argv()), 0)

            workbook = openpyxl.load_workbook(workspace.output)
            try:
                metrics = {
                    row[1]: row[2]
                    for row in workbook["resumo"].iter_rows(min_row=2, values_only=True)
                    if row[0] == "metrica"
                }
            finally:
                workbook.close()

            self.assertEqual(metrics["registros_lidos"], 2)
            self.assertEqual(metrics["registros_validos"], 1)
            self.assertEqual(metrics["registros_invalidos"], 1)

    def test_unknown_command_is_rejected(self):
        with self.assertRaises(SystemExit):
            main(["inexistente"])


class VersionTests(unittest.TestCase):
    def test_reports_the_version_and_exits_cleanly(self):
        # Quem automatiza precisa saber qual versao esta rodando antes de
        # culpar os proprios dados por uma mudanca de comportamento.
        buffer = io.StringIO()
        with redirect_stdout(buffer), self.assertRaises(SystemExit) as raised:
            main(["--versao"])

        self.assertEqual(raised.exception.code, 0)
        self.assertIn(__version__, buffer.getvalue())

    def test_the_declared_version_is_the_one_that_gets_packaged(self):
        """Duas fontes de versao divergem; esta prova que ha so uma.

        `pyproject.toml` le o atributo do pacote, entao o numero publicado no
        PyPI e o numero que a CLI imprime nao tem como discordar.
        """
        pyproject = (
            Path(__file__).resolve().parents[1] / "pyproject.toml"
        ).read_text(encoding="utf-8")

        self.assertIn(
            'version = {attr = "automacao_planilhas.__version__"}', pyproject
        )
        self.assertNotIn(f'version = "{__version__}"', pyproject)


if __name__ == "__main__":
    unittest.main()
