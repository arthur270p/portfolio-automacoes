"""Aceitação: as promessas que o README faz, verificadas de ponta a ponta."""

import io
import json
import socket
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import mock

import openpyxl

from automacao_planilhas.cli import main

PROJECT_ROOT = Path(__file__).resolve().parents[1]
EXAMPLES = PROJECT_ROOT / "exemplos"

CONFIG = {
    "colunas_obrigatorias": ["nome", "email"],
    "chaves_duplicidade": ["email"],
    "tipos": {"email": "texto"},
}


def run(input_dir: Path, output: Path, config_path: Path) -> tuple[int, str]:
    buffer = io.StringIO()
    with redirect_stdout(buffer):
        code = main(
            [
                "processar",
                "--entrada",
                str(input_dir),
                "--config",
                str(config_path),
                "--saida",
                str(output),
                "--sobrescrever",
            ]
        )
    return code, buffer.getvalue()


def workspace(root: Path, files: dict[str, str]) -> tuple[Path, Path, Path]:
    input_dir = root / "entrada"
    input_dir.mkdir()
    for name, content in files.items():
        (input_dir / name).write_text(content, encoding="utf-8")
    config_path = root / "config.json"
    config_path.write_text(json.dumps(CONFIG), encoding="utf-8")
    return input_dir, config_path, root / "relatorio.xlsx"


class OfflineTests(unittest.TestCase):
    def test_runs_with_every_network_call_blocked(self):
        """O README promete que funciona sem internet. Isto prova.

        Bloquear o socket é mais forte que procurar `import requests`: pega
        qualquer chamada de rede, inclusive a que viesse de uma dependência.
        """
        with TemporaryDirectory() as temporary:
            input_dir, config_path, output = workspace(
                Path(temporary),
                {"dados.csv": "nome,email\nAna,ana@exemplo.com.br\n"},
            )

            def refuse(*_args, **_kwargs):
                raise AssertionError("a ferramenta tentou abrir conexão de rede")

            with mock.patch.object(socket, "socket", refuse), mock.patch.object(
                socket, "create_connection", refuse
            ):
                code, _ = run(input_dir, output, config_path)

            self.assertEqual(code, 0)
            self.assertTrue(output.exists())


class DelimiterTests(unittest.TestCase):
    def test_reads_both_comma_and_semicolon_in_the_same_run(self):
        # O Excel brasileiro exporta com ponto e virgula; quase todo o resto
        # do mundo usa virgula. Uma pasta real tem os dois.
        with TemporaryDirectory() as temporary:
            input_dir, config_path, output = workspace(
                Path(temporary),
                {
                    "virgula.csv": "nome,email\nAna,ana@exemplo.com.br\n",
                    "ponto_virgula.csv": "nome;email\nBruno;bruno@exemplo.com.br\n",
                },
            )

            code, _ = run(input_dir, output, config_path)
            self.assertEqual(code, 0)

            workbook = openpyxl.load_workbook(output)
            try:
                rows = list(
                    workbook["dados_limpos"].iter_rows(min_row=2, values_only=True)
                )
                header = [cell.value for cell in workbook["dados_limpos"][1]]
            finally:
                workbook.close()

            names = {row[header.index("nome")] for row in rows}
            self.assertEqual(names, {"Ana", "Bruno"})


class MessageTests(unittest.TestCase):
    def test_speaks_portuguese_without_leaking_a_traceback(self):
        with TemporaryDirectory() as temporary:
            input_dir, config_path, output = workspace(
                Path(temporary),
                {"dados.csv": "nome,email\nAna,ana@exemplo.com.br\n"},
            )

            code, printed = run(input_dir, output, config_path)

            self.assertEqual(code, 0)
            self.assertIn("Relatório gerado", printed)
            self.assertIn("registros válidos", printed)
            self.assertNotIn("Traceback", printed)


class DocumentedExampleTests(unittest.TestCase):
    def test_the_bundled_example_matches_the_numbers_in_the_readme(self):
        """Prende os números que o README publica.

        Mexer nas fixtures sem atualizar o README passaria despercebido, e um
        portfólio que mostra um número e entrega outro perde a credibilidade
        que ele existe para construir.
        """
        with TemporaryDirectory() as temporary:
            output = Path(temporary) / "relatorio.xlsx"
            code = main(
                [
                    "processar",
                    "--entrada",
                    str(EXAMPLES / "entrada"),
                    "--config",
                    str(EXAMPLES / "config.json"),
                    "--saida",
                    str(output),
                ]
            )
            self.assertEqual(code, 0)

            workbook = openpyxl.load_workbook(output)
            try:
                self.assertEqual(
                    workbook.sheetnames, ["dados_limpos", "erros", "resumo"]
                )
                metrics = {
                    row[1]: row[2]
                    for row in workbook["resumo"].iter_rows(min_row=2, values_only=True)
                    if row[0] == "metrica"
                }
            finally:
                workbook.close()

            self.assertEqual(metrics["arquivos_encontrados"], 2)
            self.assertEqual(metrics["arquivos_processados"], 2)
            self.assertEqual(metrics["registros_lidos"], 5)
            self.assertEqual(metrics["registros_validos"], 2)
            self.assertEqual(metrics["registros_invalidos"], 3)
            self.assertEqual(metrics["duplicidades"], 1)
            self.assertEqual(metrics["problemas_encontrados"], 3)

    def test_the_readme_publishes_the_real_test_count(self):
        """Numero que envelhece sozinho perde a credibilidade que gera.

        O README vende a quantidade de testes. Acrescentar um caso sem tocar no
        texto transformaria o argumento em exagero, e nada avisaria.
        """
        readme = (PROJECT_ROOT / "README.md").read_text(encoding="utf-8")
        suite = unittest.defaultTestLoader.discover(
            str(PROJECT_ROOT / "tests"), top_level_dir=str(PROJECT_ROOT)
        )

        self.assertIn(f"{suite.countTestCases()} testes", readme)


class AliasAcceptanceTests(unittest.TestCase):
    def test_consolidates_files_that_disagree_on_the_column_name(self):
        """O caso que motiva a funcionalidade, de ponta a ponta.

        Dois sistemas, dois nomes para a mesma coluna. Sem apelido, um dos
        arquivos seria recusado inteiro e a consolidacao voltaria a ser manual.
        """
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_dir = root / "entrada"
            input_dir.mkdir()
            (input_dir / "sistema_a.csv").write_text(
                "nome,email\nAna,ana@exemplo.com.br\n", encoding="utf-8"
            )
            (input_dir / "sistema_b.csv").write_text(
                "Nome;E-mail do Cliente\nBruno;bruno@exemplo.com.br\n",
                encoding="utf-8",
            )
            config_path = root / "config.json"
            config_path.write_text(
                json.dumps({**CONFIG, "apelidos": {"email": ["e_mail_do_cliente"]}}),
                encoding="utf-8",
            )
            output = root / "relatorio.xlsx"

            code, _ = run(input_dir, output, config_path)
            self.assertEqual(code, 0)

            workbook = openpyxl.load_workbook(output)
            try:
                header = [cell.value for cell in workbook["dados_limpos"][1]]
                rows = list(
                    workbook["dados_limpos"].iter_rows(min_row=2, values_only=True)
                )
            finally:
                workbook.close()

            emails = {row[header.index("email")] for row in rows}
            self.assertEqual(
                emails, {"ana@exemplo.com.br", "bruno@exemplo.com.br"}
            )



if __name__ == "__main__":
    unittest.main()
